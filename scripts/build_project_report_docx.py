from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
DOCX_PATH = DOCS / "AIpindou项目分析报告.docx"
MD_PATH = DOCS / "AIpindou项目分析报告.md"


TITLE = "AIpindou 项目分析报告"
SUBTITLE = "图片转拼豆图纸桌面软件接手评估"
DATE_LINE = "生成日期：2026-05-28"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_grid = tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        tbl.insert(0, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9.5)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, "F2F4F7")
        set_cell_text(cell, header, bold=True)

    for row_values in rows:
        row = table.add_row()
        for index, value in enumerate(row_values):
            set_cell_text(row.cells[index], value)

    document.add_paragraph()


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_para(document: Document, text: str, style: str | None = None, bold_prefix: str | None = None) -> None:
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(6)
    if bold_prefix and text.startswith(bold_prefix):
        prefix = paragraph.add_run(bold_prefix)
        prefix.bold = True
        prefix.font.name = "Microsoft YaHei"
        prefix._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        rest = paragraph.add_run(text[len(bold_prefix):])
        rest.font.name = "Microsoft YaHei"
        rest._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    else:
        run = paragraph.add_run(text)
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        add_para(document, item, style="List Bullet")


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size, color in [
        ("Heading 1", 16, "2E74B5"),
        ("Heading 2", 13, "2E74B5"),
        ("Heading 3", 12, "1F4D78"),
    ]:
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10 if style_name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(5)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("AIpindou 项目分析报告")


def add_cover(document: Document) -> None:
    document.add_paragraph()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(TITLE)
    run.bold = True
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor.from_string("0B2545")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(SUBTITLE)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string("555555")

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(DATE_LINE)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(10)

    document.add_paragraph()
    add_para(
        document,
        "本报告基于当前仓库的只读分析生成，用于项目接手、技术判断和后续迭代规划。结论聚焦可运行性、功能完成度、风险点和下一步交付路线。",
    )
    document.add_section(WD_SECTION.NEW_PAGE)


def build_docx() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    document = Document()
    configure_document(document)
    add_cover(document)

    add_heading(document, "一、执行摘要", 1)
    add_para(
        document,
        "AIpindou 当前已经具备较完整的前端闭环，不是需要从零开发的项目。它已经覆盖图片导入、网格化、色卡匹配、预览统计、局部微调和 PNG/JSON/CSV 导出等核心流程。",
    )
    add_para(
        document,
        "总体判断：前端 MVP 已基本成型，部分增强能力已经提前实现；桌面壳具备打包配置；后端目录更像历史或预留代码，当前不应作为 MVP 运行依赖。",
    )

    add_heading(document, "二、项目概览", 1)
    add_para(document, "项目定位：面向拼豆创作用户的桌面软件，把图片自动转换为可制作的拼豆图纸。")
    add_para(document, "当前最小闭环：导入图片、按网格生成拼豆图、匹配基础色卡、预览统计、导出 PNG/JSON，并打包为 Windows exe。")
    add_para(document, "当前技术路线：Vue 3 + TypeScript + Vite + Element Plus + Electron，图像处理主要通过浏览器 Canvas 和本地模型完成。")

    add_table(
        document,
        ["维度", "当前状态", "评估"],
        [
            ["前端", "Vue 3 单页应用，核心页面集中在 Home.vue", "主流程清晰，功能集中，便于继续迭代"],
            ["桌面端", "Electron 主进程、preload、Provider 存储和打包配置已存在", "已具备 Windows portable exe 打包基础"],
            ["图像算法", "Canvas 网格化、主色/平均色采样、近似色合并", "可满足基础拼豆图生成"],
            ["AI 能力", "本地背景分割、面部轮廓增强、外部 AI Provider", "亮点较多，但需补隐私提示和稳定性验证"],
            ["后端", "Express 代码存在但 service 缺失", "不建议纳入当前 MVP 依赖"],
        ],
        [1500, 3900, 3960],
    )

    add_heading(document, "三、技术架构分析", 1)
    add_heading(document, "3.1 前端架构", 2)
    add_para(document, "前端位于 frontend/，采用 Vue 3、TypeScript、Vite、Element Plus。核心页面为 frontend/src/views/Home.vue。")
    add_bullets(
        document,
        [
            "pixelation.ts：负责图片到拼豆网格的主算法。",
            "pattern-render.ts：负责画布渲染、网格线、色号和导出图例。",
            "pattern-export.ts：负责 JSON 和 CSV 数据导出。",
            "pixel-editing.ts：负责排除颜色、替换单格颜色、刷新统计。",
            "background-segmentation.ts：负责本地背景分割流程。",
            "face-contour-enhancement.ts：负责面部轮廓增强。",
            "watermark.ts：负责明水印和 DCT 暗水印。",
        ],
    )

    add_heading(document, "3.2 Electron 桌面壳", 2)
    add_para(document, "Electron 代码位于 electron/，负责桌面窗口、IPC 暴露、自定义 AI Provider 配置存储和外部图像模型调用。")
    add_para(document, "根 package.json 已配置 electron-builder，目标为 Windows portable 包，产物目录为 release。")

    add_heading(document, "3.3 后端目录", 2)
    add_para(document, "backend/ 中存在 Express controller，但引用的 ImageService、ColorService、ExportService 未在仓库中出现。当前应视为预留或历史代码，不建议作为运行链路的一部分。")

    add_heading(document, "四、已实现功能", 1)
    add_table(
        document,
        ["功能模块", "已实现内容", "成熟度"],
        [
            ["图片导入", "点击选择、拖拽导入，支持 PNG/JPEG/WebP/GIF", "较完整"],
            ["网格生成", "横向颗数可调，纵向按原图比例自动计算", "较完整"],
            ["转换算法", "主色模式、平均色模式、近似色合并、透明格处理", "可用"],
            ["色卡系统", "MARD 多档位色卡，多厂商编号映射", "较强"],
            ["预览统计", "圆珠/方格渲染、总颗数、颜色数、用色清单", "较完整"],
            ["微调编辑", "画笔、吸管、缩放、单格改色、辅助显示", "基础可用"],
            ["背景擦除", "ONNX U2NetP，本地 WebGPU/WASM 推理", "增强能力"],
            ["轮廓增强", "MediaPipe face landmarker，启发式兜底", "增强能力"],
            ["导出", "PNG、JSON、CSV", "缺 PDF/分页打印"],
            ["AI 优化", "自定义 Provider，DashScope Wan2.6 预设", "需安全提示"],
        ],
        [1700, 5600, 2060],
    )

    add_heading(document, "五、与目标需求匹配度", 1)
    add_table(
        document,
        ["目标需求", "当前匹配度", "说明"],
        [
            ["自动生成拼豆图纸", "高", "核心转换链路已经打通"],
            ["风格调整", "中", "已有主色/平均、渲染样式、AI 优化，但缺少产品化风格预设"],
            ["自定义色块调整", "中", "已有单格改色和排除颜色，缺批量替换、撤销/重做、自定义色卡"],
            ["色卡匹配", "高", "MARD 档位和多厂商编号已实现"],
            ["导出图纸", "中", "PNG/JSON/CSV 已实现，PDF 和打印布局缺失"],
            ["Windows exe", "中", "打包配置存在，但依赖未安装，尚未验证"],
        ],
        [2200, 1700, 5460],
    )

    add_heading(document, "六、主要风险与问题", 1)
    add_bullets(
        document,
        [
            "依赖未安装：当前 node_modules 和 frontend/node_modules 不存在，npm test 因 vitest 缺失而失败。",
            "后端不可用：backend controller 引用缺失 service，短期不应纳入交付链路。",
            "旧文档与实际技术栈不一致：旧文档仍出现 React、Ant Design、Redux、Konva 等描述，需清理。",
            "PDF/打印能力缺失：现有 PNG 导出可用，但真实制作场景更需要 A4 分页和 PDF 图纸。",
            "编辑效率不足：单格微调已可用，但缺批量替换、区域编辑、撤销/重做和工程保存。",
            "隐私提示不足：本地 AI 与外部 AI Provider 的边界需要在 UI 和文档中明确。",
            "色卡数据需专项校验：多厂商编号和 RGB 数据直接影响真实采购和制作准确性。",
        ],
    )

    add_heading(document, "七、产品短板", 1)
    add_para(document, "当前项目的技术能力已经超过普通 MVP，但产品化体验仍有几处短板。")
    add_bullets(
        document,
        [
            "缺少工程文件保存和再次打开能力，用户难以长期编辑同一图纸。",
            "缺少面向普通用户的风格预设，参数虽多但理解成本偏高。",
            "缺少打印友好的 PDF 图纸和分页策略。",
            "缺少批量编辑能力，真实大图微调成本较高。",
            "缺少新手引导和隐私说明，特别是外部 AI 调用场景。",
        ],
    )

    add_heading(document, "八、建议迭代路线", 1)
    add_table(
        document,
        ["阶段", "目标", "关键任务"],
        [
            ["第一阶段", "稳定接手", "安装依赖；跑通测试、构建和 exe 打包；修正文档；标记后端为预留"],
            ["第二阶段", "完成图纸交付闭环", "增加 PDF 导出、A4 分页、标题/尺寸/用色清单和工程保存"],
            ["第三阶段", "强化编辑体验", "撤销/重做、批量颜色替换、区域选择、原图对照、色号搜索"],
            ["第四阶段", "产品化风格能力", "头像清晰、卡通高饱和、照片柔和、少色省豆、轮廓增强等预设"],
        ],
        [1500, 2500, 5360],
    )

    add_heading(document, "九、结论建议", 1)
    add_para(document, "AIpindou 已经有一个可继续推进的好底子。下一步不建议重写，而应优先验证可运行性、收敛产品边界并补齐真实交付所需的 PDF/打印与编辑工作流。")
    add_para(document, "建议近期目标是把项目从“技术原型”推进到“可给真实拼豆用户试用”的版本。最关键的三件事是：")
    add_bullets(
        document,
        [
            "验证工程可运行、可测试、可打包。",
            "补齐 PDF、分页打印和工程保存能力。",
            "强化批量编辑、自定义色块和风格预设体验。",
        ],
    )

    add_heading(document, "附录：手机阅读交付方式", 1)
    add_bullets(
        document,
        [
            "飞书云文档：推荐方式，新建后可在手机飞书中直接打开阅读。",
            "本地 PDF：保存在项目 docs 目录，可通过微信文件传输助手、飞书文件、数据线或网盘同步到手机。",
            "邮件、微信、短信直发：当前环境没有已配置连接器，不能直接发送。",
            "二维码：如果飞书返回可访问链接，后续可生成二维码图片用于手机扫码。",
        ],
    )

    document.save(DOCX_PATH)
    write_markdown()


def write_markdown() -> None:
    content = f"""# {TITLE}

{SUBTITLE}

{DATE_LINE}

## 一、执行摘要

AIpindou 当前已经具备较完整的前端闭环，不是需要从零开发的项目。它已经覆盖图片导入、网格化、色卡匹配、预览统计、局部微调和 PNG/JSON/CSV 导出等核心流程。

总体判断：前端 MVP 已基本成型，部分增强能力已经提前实现；桌面壳具备打包配置；后端目录更像历史或预留代码，当前不应作为 MVP 运行依赖。

## 二、项目概览

- 项目定位：面向拼豆创作用户的桌面软件，把图片自动转换为可制作的拼豆图纸。
- 当前最小闭环：导入图片、按网格生成拼豆图、匹配基础色卡、预览统计、导出 PNG/JSON，并打包为 Windows exe。
- 当前技术路线：Vue 3 + TypeScript + Vite + Element Plus + Electron，图像处理主要通过浏览器 Canvas 和本地模型完成。

## 三、技术架构分析

### 前端架构

前端位于 `frontend/`，采用 Vue 3、TypeScript、Vite、Element Plus。核心页面为 `frontend/src/views/Home.vue`。

- `pixelation.ts`：负责图片到拼豆网格的主算法。
- `pattern-render.ts`：负责画布渲染、网格线、色号和导出图例。
- `pattern-export.ts`：负责 JSON 和 CSV 数据导出。
- `pixel-editing.ts`：负责排除颜色、替换单格颜色、刷新统计。
- `background-segmentation.ts`：负责本地背景分割流程。
- `face-contour-enhancement.ts`：负责面部轮廓增强。
- `watermark.ts`：负责明水印和 DCT 暗水印。

### Electron 桌面壳

Electron 代码位于 `electron/`，负责桌面窗口、IPC 暴露、自定义 AI Provider 配置存储和外部图像模型调用。根 `package.json` 已配置 electron-builder，目标为 Windows portable 包，产物目录为 `release`。

### 后端目录

`backend/` 中存在 Express controller，但引用的 `ImageService`、`ColorService`、`ExportService` 未在仓库中出现。当前应视为预留或历史代码，不建议作为运行链路的一部分。

## 四、已实现功能

- 图片导入：点击选择、拖拽导入，支持 PNG/JPEG/WebP/GIF。
- 网格生成：横向颗数可调，纵向按原图比例自动计算。
- 转换算法：主色模式、平均色模式、近似色合并、透明格处理。
- 色卡系统：MARD 多档位色卡，多厂商编号映射。
- 预览统计：圆珠/方格渲染、总颗数、颜色数、用色清单。
- 微调编辑：画笔、吸管、缩放、单格改色、辅助显示。
- 背景擦除：ONNX U2NetP，本地 WebGPU/WASM 推理。
- 轮廓增强：MediaPipe face landmarker，启发式兜底。
- 导出：PNG、JSON、CSV；仍缺 PDF/分页打印。
- AI 优化：自定义 Provider，DashScope Wan2.6 预设。

## 五、与目标需求匹配度

| 目标需求 | 当前匹配度 | 说明 |
| --- | --- | --- |
| 自动生成拼豆图纸 | 高 | 核心转换链路已经打通 |
| 风格调整 | 中 | 已有主色/平均、渲染样式、AI 优化，但缺少产品化风格预设 |
| 自定义色块调整 | 中 | 已有单格改色和排除颜色，缺批量替换、撤销/重做、自定义色卡 |
| 色卡匹配 | 高 | MARD 档位和多厂商编号已实现 |
| 导出图纸 | 中 | PNG/JSON/CSV 已实现，PDF 和打印布局缺失 |
| Windows exe | 中 | 打包配置存在，但依赖未安装，尚未验证 |

## 六、主要风险与问题

- 依赖未安装：当前 `node_modules` 和 `frontend/node_modules` 不存在，`npm test` 因 `vitest` 缺失而失败。
- 后端不可用：backend controller 引用缺失 service，短期不应纳入交付链路。
- 旧文档与实际技术栈不一致：旧文档仍出现 React、Ant Design、Redux、Konva 等描述，需清理。
- PDF/打印能力缺失：现有 PNG 导出可用，但真实制作场景更需要 A4 分页和 PDF 图纸。
- 编辑效率不足：单格微调已可用，但缺批量替换、区域编辑、撤销/重做和工程保存。
- 隐私提示不足：本地 AI 与外部 AI Provider 的边界需要在 UI 和文档中明确。
- 色卡数据需专项校验：多厂商编号和 RGB 数据直接影响真实采购和制作准确性。

## 七、产品短板

- 缺少工程文件保存和再次打开能力，用户难以长期编辑同一图纸。
- 缺少面向普通用户的风格预设，参数虽多但理解成本偏高。
- 缺少打印友好的 PDF 图纸和分页策略。
- 缺少批量编辑能力，真实大图微调成本较高。
- 缺少新手引导和隐私说明，特别是外部 AI 调用场景。

## 八、建议迭代路线

| 阶段 | 目标 | 关键任务 |
| --- | --- | --- |
| 第一阶段 | 稳定接手 | 安装依赖；跑通测试、构建和 exe 打包；修正文档；标记后端为预留 |
| 第二阶段 | 完成图纸交付闭环 | 增加 PDF 导出、A4 分页、标题/尺寸/用色清单和工程保存 |
| 第三阶段 | 强化编辑体验 | 撤销/重做、批量颜色替换、区域选择、原图对照、色号搜索 |
| 第四阶段 | 产品化风格能力 | 头像清晰、卡通高饱和、照片柔和、少色省豆、轮廓增强等预设 |

## 九、结论建议

AIpindou 已经有一个可继续推进的好底子。下一步不建议重写，而应优先验证可运行性、收敛产品边界并补齐真实交付所需的 PDF/打印与编辑工作流。

建议近期目标是把项目从“技术原型”推进到“可给真实拼豆用户试用”的版本。最关键的三件事是：

- 验证工程可运行、可测试、可打包。
- 补齐 PDF、分页打印和工程保存能力。
- 强化批量编辑、自定义色块和风格预设体验。

## 附录：手机阅读交付方式

- 飞书云文档：推荐方式，新建后可在手机飞书中直接打开阅读。
- 本地 PDF：保存在项目 docs 目录，可通过微信文件传输助手、飞书文件、数据线或网盘同步到手机。
- 邮件、微信、短信直发：当前环境没有已配置连接器，不能直接发送。
- 二维码：如果飞书返回可访问链接，后续可生成二维码图片用于手机扫码。
"""
    MD_PATH.write_text(content, encoding="utf-8")


if __name__ == "__main__":
    build_docx()
    print(DOCX_PATH)
    print(MD_PATH)
